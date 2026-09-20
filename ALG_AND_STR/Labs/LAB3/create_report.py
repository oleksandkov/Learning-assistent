"""
Генератор звіту у форматах DOCX та PDF для Лабораторної роботи №3, Варіант 12.
Базується безпосередньо на Звіт_лабораторна1.docx для 100% збереження титульної сторінки,
шрифтів, стилів, колонтитулів та полів документа.
"""
import os
import re
import docx
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import win32com.client

def set_para_shading(p, color_hex):
    shading_xml = f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="{color_hex}"/>'
    p._element.get_or_add_pPr().append(parse_xml(shading_xml))

def add_code_line(p, line):
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    set_para_shading(p, "1F1F1F")

    tokens = re.split(r'("(?:\\.|[^"\\])*"|\b(?:include|struct|constexpr|int|char|void|bool|nullptr|if|else|while|for|switch|case|break|return|new|delete|const|static|auto|true|false|assert|string|vector)\b)', line)
    keywords = {"include", "struct", "constexpr", "int", "char", "void", "bool", "nullptr", "if", "else", "while", "for", "switch", "case", "break", "return", "new", "delete", "const", "static", "auto", "true", "false", "assert"}
    types = {"string", "vector"}

    for token in tokens:
        if not token:
            continue
        r = p.add_run(token)
        r.font.name = "Consolas"
        r.font.size = Pt(9.5)

        if token.startswith('"'):
            r.font.color.rgb = RGBColor(0xCE, 0x91, 0x78)
        elif token.startswith("#") or token in keywords:
            r.font.color.rgb = RGBColor(0x56, 0x9C, 0xD6) if token not in {"include", "return", "if", "else", "while", "for", "switch", "case", "break"} else RGBColor(0xC5, 0x86, 0xC0)
        elif token in types:
            r.font.color.rgb = RGBColor(0x4E, 0xC9, 0xB0)
        else:
            r.font.color.rgb = RGBColor(0xDC, 0xDC, 0xDC)

def build_report():
    base_dir = r"C:\Github\Learning-assistent\ALG_AND_STR\Labs\LAB3"
    template_docx = r"C:\Github\Learning-assistent\ALG_AND_STR\Labs\LAB1\v2\Звіт_лабораторна1.docx"
    docx_path = os.path.join(base_dir, "Звіт_лабораторна3.docx")
    pdf_path = os.path.join(base_dir, "Звіт_лабораторна3.pdf")
    flowchart_img = os.path.join(base_dir, "flowchart.png")
    code_path = os.path.join(base_dir, "main.cpp")

    with open(code_path, "r", encoding="utf-8") as f:
        code_lines = f.readlines()

    # Відкриваємо оригінальний файл для ідеального збереження структури
    doc = Document(template_docx)

    # 1. Оновлення титульної сторінки (P0..P19)
    # Зберігаємо точну геометрію, розміри, шрифти та відступи
    doc.paragraphs[6].runs[0].text = "Дослідження та реалізація базових динамічних структур даних"
    doc.paragraphs[8].runs[0].text = "Звіт про виконання лабораторної роботи № 3"

    # 2. Видаляємо старі параграфи лаби 1 (починаючи з P20)
    for p in list(doc.paragraphs)[20:]:
        p._element.getparent().remove(p._element)

    # 3. Додаємо зміст Лабораторної роботи №3
    # Тема
    p_theme = doc.add_paragraph()
    p_theme.paragraph_format.space_before = Pt(6)
    p_theme.paragraph_format.space_after = Pt(4)
    r_th1 = p_theme.add_run("Тема: ")
    r_th1.font.name = "Times New Roman"
    r_th1.font.size = Pt(16)
    r_th1.bold = True
    r_th2 = p_theme.add_run("Дослідження та реалізація базових структур даних: двозв'язний список з розгалуженнями. Варіант 12.")
    r_th2.font.name = "Times New Roman"
    r_th2.font.size = Pt(14)

    # Мета
    p_goal = doc.add_paragraph()
    p_goal.paragraph_format.space_before = Pt(4)
    p_goal.paragraph_format.space_after = Pt(12)
    r_g1 = p_goal.add_run("Мета: ")
    r_g1.font.name = "Times New Roman"
    r_g1.font.size = Pt(16)
    r_g1.bold = True
    r_g2 = p_goal.add_run("дослідити та реалізувати базові динамічні структури даних; спроєктувати та реалізувати на мові C++ двозв‘язний список із розгалуженнями для збереження та обробки послідовностей чисел відповідно до варіанту 12; набути навичок динамічного керування пам'яттю та роботи з покажчиками.")
    r_g2.font.name = "Times New Roman"
    r_g2.font.size = Pt(14)

    # Завдання
    p_task_title = doc.add_paragraph()
    p_task_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_task_title.paragraph_format.space_before = Pt(10)
    p_task_title.paragraph_format.space_after = Pt(6)
    r_tt = p_task_title.add_run("ЗАВДАННЯ")
    r_tt.font.name = "Times New Roman"
    r_tt.font.size = Pt(16)
    r_tt.bold = True

    p_task = doc.add_paragraph()
    p_task.paragraph_format.space_before = Pt(2)
    p_task.paragraph_format.space_after = Pt(12)
    p_task.paragraph_format.line_spacing = 1.15
    r_t = p_task.add_run(
        "Реалізувати двозв‘язний список з розгалуженнями для зберігання і операцій з даними виду:\n"
        "Ім'я  |  Послідовність чисел\n\n"
        "Забезпечити виконання операцій:\n"
        "• додавання елементів в список;\n"
        "• пошук і роздрукування елементу в списку;\n"
        "• видалення елементу зі списку;\n"
        "• підрахунок елементів у списку;\n"
        "• роздрукування списку.\n\n"
        "Примітка. Послідовність чисел може містити від 1 до N чисел. Для послідовностей завдовжки більш K (K < N) організувати гілки. У роботі прийнято: K = 3, N = 10."
    )
    r_t.font.name = "Times New Roman"
    r_t.font.size = Pt(14)

    # Код програми
    p_code_title = doc.add_paragraph()
    p_code_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_code_title.paragraph_format.space_before = Pt(10)
    p_code_title.paragraph_format.space_after = Pt(6)
    r_ct = p_code_title.add_run("КОД ПРОГРАМИ")
    r_ct.font.name = "Times New Roman"
    r_ct.font.size = Pt(16)
    r_ct.bold = True

    for line in code_lines:
        p_l = doc.add_paragraph()
        add_code_line(p_l, line.rstrip("\n"))

    # Результати тестування
    doc.add_page_break()
    p_res_title = doc.add_paragraph()
    p_res_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_res_title.paragraph_format.space_before = Pt(10)
    p_res_title.paragraph_format.space_after = Pt(6)
    r_rt = p_res_title.add_run("РЕЗУЛЬТАТИ ТЕСТУВАННЯ ТА РОБОТИ ПРОГРАМИ")
    r_rt.font.name = "Times New Roman"
    r_rt.font.size = Pt(16)
    r_rt.bold = True

    p_test_desc = doc.add_paragraph()
    p_test_desc.paragraph_format.space_before = Pt(2)
    p_test_desc.paragraph_format.space_after = Pt(4)
    r_td = p_test_desc.add_run("1. Результати виконання вбудованих автоматичних тестів (assert) при запуску з прапорцем --test:")
    r_td.font.name = "Times New Roman"
    r_td.font.size = Pt(14)

    test_logs = [
        "> g++ -std=c++20 -O2 main.cpp -o main.exe",
        "> ./main.exe --test",
        "Автотести успішно пройдено."
    ]
    for tl in test_logs:
        p_tl = doc.add_paragraph()
        p_tl.paragraph_format.space_before = Pt(0)
        p_tl.paragraph_format.space_after = Pt(0)
        set_para_shading(p_tl, "1F1F1F")
        r = p_tl.add_run(tl)
        r.font.name = "Consolas"
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(0x4E, 0xC9, 0xB0) if "успішно" in tl else RGBColor(0xDC, 0xDC, 0xDC)

    p_run_desc = doc.add_paragraph()
    p_run_desc.paragraph_format.space_before = Pt(8)
    p_run_desc.paragraph_format.space_after = Pt(4)
    r_rd = p_run_desc.add_run("2. Протокол виконання демонстраційного сценарію в інтерактивному режимі:")
    r_rd.font.name = "Times New Roman"
    r_rd.font.size = Pt(14)

    demo_logs = [
        "=== Лабораторна робота №3. Варіант 12 ===",
        "Двозв'язний список із розгалуженнями (K = 3, N = 10)",
        "",
        "Меню операцій: [6 - Завантажити зразки, 5 - Друк, 2 - Пошук, 3 - Видалення, 4 - Кількість]",
        "Ваш вибір: 6",
        "Завантажено 4 зразкові записи (із гілками та без).",
        "",
        "=== Вміст списку (всього елементів: 4) ===",
        "1) Елемент: Запис_1",
        "   Основний вузол (<= 3 чисел): [5, 12]",
        "   Гілка розгалуження (> 3 чисел): відсутня",
        "2) Елемент: Запис_2",
        "   Основний вузол (<= 3 чисел): [10, 20, 30]",
        "   Гілка розгалуження (> 3 чисел): [40 -> 50 -> 60]",
        "3) Елемент: Запис_3",
        "   Основний вузол (<= 3 чисел): [7, 8, 9]",
        "   Гілка розгалуження (> 3 чисел): відсутня",
        "4) Елемент: Запис_4",
        "   Основний вузол (<= 3 чисел): [100, 200, 300]",
        "   Гілка розгалуження (> 3 чисел): [400]",
        "==========================================",
        "",
        "Пошук елемента 'Запис_2':",
        "Елемент: Запис_2",
        "   Основний вузол (<= 3 чисел): [10, 20, 30]",
        "   Гілка розгалуження (> 3 чисел): [40 -> 50 -> 60]",
        "",
        "Видалення елемента 'Запис_1': Елемент \"Запис_1\" видалено.",
        "Кількість елементів у списку: 3"
    ]
    for dl in demo_logs:
        p_dl = doc.add_paragraph()
        p_dl.paragraph_format.space_before = Pt(0)
        p_dl.paragraph_format.space_after = Pt(0)
        set_para_shading(p_dl, "1F1F1F")
        r = p_dl.add_run(dl)
        r.font.name = "Consolas"
        r.font.size = Pt(9.5)
        if dl.startswith("1)") or dl.startswith("2)") or dl.startswith("3)") or dl.startswith("4)"):
            r.font.color.rgb = RGBColor(0x56, 0x9C, 0xD6)
        elif "Гілка" in dl:
            r.font.color.rgb = RGBColor(0xCE, 0x91, 0x78)
        else:
            r.font.color.rgb = RGBColor(0xDC, 0xDC, 0xDC)

    # Блок-схема (окрема сторінка)
    doc.add_page_break()
    p_flow_title = doc.add_paragraph()
    p_flow_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_flow_title.paragraph_format.space_before = Pt(4)
    p_flow_title.paragraph_format.space_after = Pt(6)
    r_ft = p_flow_title.add_run("БЛОК-СХЕМА АЛГОРИТМУ (UML)")
    r_ft.font.name = "Times New Roman"
    r_ft.font.size = Pt(16)
    r_ft.bold = True

    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(2)
    p_img.paragraph_format.space_after = Pt(2)
    p_img.add_run("[SVG_FLOWCHART_PLACEHOLDER]")

    # Висновок (окрема сторінка)
    doc.add_page_break()
    p_conc = doc.add_paragraph()
    p_conc.paragraph_format.space_before = Pt(12)
    p_conc.paragraph_format.space_after = Pt(6)
    p_conc.paragraph_format.line_spacing = 1.15
    r_c_b = p_conc.add_run("Висновок: ")
    r_c_b.font.name = "Times New Roman"
    r_c_b.font.size = Pt(16)
    r_c_b.bold = True

    r_c_text = p_conc.add_run(
        "У результаті виконання лабораторної роботи було успішно досліджено та реалізовано динамічну структуру даних — "
        "двозв‘язний список із розгалуженнями для опрацювання числових послідовностей змінної довжини (варіант 12). "
        "Розроблена консольна програма на мові C++ реалізує ефективний розподіл елементів: послідовності довжиною до K = 3 "
        "чисел зберігаються безпосередньо у полях основного вузла списку, а числа понад поріг K динамічно виносяться в однозв'язні "
        "гілки-підсписки. Повністю забезпечено виконання регламентованих операцій: додавання, пошуку, видалення елементів "
        "із коректним вивільненням пам'яті гілок, підрахунку розміру списку за O(1) та наочного роздрукування всієї ієрархічної структури. "
        "Складено алгоритм, блок-схему UML та здійснено вичерпне автоматизоване й інтерактивне тестування."
    )
    r_c_text.font.name = "Times New Roman"
    r_c_text.font.size = Pt(14)

    doc.save(docx_path)
    print(f"Попередній DOCX збережено: {docx_path}")

    # Вставка SVG безпосередньо через Word COM Automation та експорт у PDF
    svg_file = os.path.abspath(os.path.join(base_dir, "flowchart.svg"))
    try:
        word = win32com.client.DispatchEx("Word.Application")
        word.Visible = False
        doc_word = word.Documents.Open(docx_path)
        
        find_range = doc_word.Content
        if find_range.Find.Execute(FindText="[SVG_FLOWCHART_PLACEHOLDER]"):
            find_range.Text = ""
            shape = find_range.InlineShapes.AddPicture(FileName=svg_file)
            shape.LockAspectRatio = -1 # msoTrue
            shape.Width = 16.0 * 28.3465 # 16.0 cm in points (fills page width)
            shape.Height = 23.0 * 28.3465 # 23.0 cm in points (fills page height)
            print(f"SVG успішно вбудовано у DOCX: {svg_file} (висота: {shape.Height} pt, ширина: {shape.Width} pt)")
        else:
            print("Попередження: маркер [SVG_FLOWCHART_PLACEHOLDER] не знайдено!")

        doc_word.Save()
        print(f"DOCX з вбудованим SVG збережено: {docx_path}")

        doc_word.SaveAs(pdf_path, FileFormat=17) # 17 = wdFormatPDF
        doc_word.Close()
        word.Quit()
        print(f"PDF створено успішно через Word: {pdf_path}")
    except Exception as e:
        print(f"Помилка Word COM: {e}")

if __name__ == "__main__":
    build_report()
