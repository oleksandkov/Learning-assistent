"""
Генератор звіту у форматах DOCX та PDF для Лабораторної роботи №4, Варіант 12.
Базується на структурі та титульній сторінці Лабораторної №1 / №3.
Вбудовує чистий C++ код без коментарів, точні результати бенчмарку та векторну блок-схему SVG.
"""
import os
import re
import docx
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import win32com.client

def set_para_shading(p, color_hex):
    shading_xml = f'<w:shd {nsdecls("w")} w:val="clear" w:color="auto" w:fill="{color_hex}"/>'
    p._element.get_or_add_pPr().append(parse_xml(shading_xml))

def add_code_line(p, line):
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    set_para_shading(p, "1F1F1F")

    tokens = re.split(r'("(?:\\.|[^"\\])*"|\b(?:include|struct|enum|class|typename|constexpr|size_t|int|char|void|bool|nullptr|if|else|while|for|switch|case|break|return|new|delete|const|static|auto|true|false|assert|string|vector|chrono|steady_clock|mt19937|uniform_int_distribution)\b)', line)
    keywords = {"include", "struct", "enum", "class", "typename", "constexpr", "size_t", "int", "char", "void", "bool", "nullptr", "if", "else", "while", "for", "switch", "case", "break", "return", "new", "delete", "const", "static", "auto", "true", "false", "assert"}
    types = {"string", "vector", "chrono", "steady_clock", "mt19937", "uniform_int_distribution"}

    for token in tokens:
        if not token:
            continue
        r = p.add_run(token)
        r.font.name = "Consolas"
        r.font.size = Pt(9.5)

        if token.startswith('"'):
            r.font.color.rgb = RGBColor(0xCE, 0x91, 0x78)
        elif token.startswith("#") or token in keywords:
            r.font.color.rgb = RGBColor(0x56, 0x9C, 0xD6) if token not in {"include", "return", "if", "else", "while", "for", "switch", "case", "break"} else RGBColor(0xC5, 0x86, 0xC0)
        elif token in types:
            r.font.color.rgb = RGBColor(0x4E, 0xC9, 0xB0)
        else:
            r.font.color.rgb = RGBColor(0xDC, 0xDC, 0xDC)

def build_report():
    base_dir = r"C:\Github\Learning-assistent\ALG_AND_STR\Labs\LAB4"
    template_docx = r"C:\Github\Learning-assistent\ALG_AND_STR\Labs\LAB1\v2\Звіт_лабораторна1.docx"
    docx_path = os.path.join(base_dir, "Звіт_лабораторна4.docx")
    pdf_path = os.path.join(base_dir, "Звіт_лабораторна4.pdf")
    svg_file = os.path.abspath(os.path.join(base_dir, "flowchart.svg"))
    code_path = os.path.join(base_dir, "main.cpp")

    with open(code_path, "r", encoding="utf-8") as f:
        code_lines = [l.rstrip("\r\n") for l in f.readlines()]

    doc = Document(template_docx)

    # 1. Титульна сторінка
    doc.paragraphs[6].runs[0].text = "Алгоритми сортування: сортування вставкою та двійковим деревом"
    doc.paragraphs[8].runs[0].text = "Звіт про виконання лабораторної роботи № 4"

    # 2. Видалення старих параграфів після титулки
    for p in list(doc.paragraphs)[20:]:
        p._element.getparent().remove(p._element)

    # 3. Тема
    p_theme = doc.add_paragraph()
    p_theme.paragraph_format.space_before = Pt(6)
    p_theme.paragraph_format.space_after = Pt(4)
    r_th1 = p_theme.add_run("Тема: ")
    r_th1.font.name = "Times New Roman"
    r_th1.font.size = Pt(16)
    r_th1.bold = True
    r_th2 = p_theme.add_run("Дослідження та порівняльний аналіз алгоритмів сортування: сортування вставкою (Insertion sort) та сортування двійковим деревом (Tree sort). Варіант 12.")
    r_th2.font.name = "Times New Roman"
    r_th2.font.size = Pt(14)

    # Мета
    p_goal = doc.add_paragraph()
    p_goal.paragraph_format.space_before = Pt(4)
    p_goal.paragraph_format.space_after = Pt(12)
    r_g1 = p_goal.add_run("Мета: ")
    r_g1.font.name = "Times New Roman"
    r_g1.font.size = Pt(16)
    r_g1.bold = True
    r_g2 = p_goal.add_run("ознайомитися на практиці з відомими алгоритмами сортування, теоретично оцінити та порівняти між собою час виконання двох різних алгоритмів (сортування вставкою та сортування за допомогою двійкового дерева), експериментально перевірити отримані оцінки на масивах різної розмірності та структури вхідних даних.")
    r_g2.font.name = "Times New Roman"
    r_g2.font.size = Pt(14)

    # Завдання
    p_task_title = doc.add_paragraph()
    p_task_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_task_title.paragraph_format.space_before = Pt(10)
    p_task_title.paragraph_format.space_after = Pt(6)
    r_tt = p_task_title.add_run("ЗАВДАННЯ")
    r_tt.font.name = "Times New Roman"
    r_tt.font.size = Pt(16)
    r_tt.bold = True

    p_task = doc.add_paragraph()
    p_task.paragraph_format.space_before = Pt(2)
    p_task.paragraph_format.space_after = Pt(12)
    p_task.paragraph_format.line_spacing = 1.15
    r_t = p_task.add_run(
        "Відповідно до варіанту 12 (алгоритми 3 та 8) виконати:\n"
        "1. Програмно реалізувати на мові C++ алгоритм сортування вставкою (Insertion sort);\n"
        "2. Програмно реалізувати алгоритм сортування за допомогою двійкового дерева пошуку (Tree sort);\n"
        "3. Провести експериментальне вимірювання часу сортування масивів різної розмірності (від 1 000 до 100 000 елементів) для різних розподілів даних: випадкові, вже відсортовані та зворотньо відсортовані;\n"
        "4. Побудувати зведену порівняльну таблицю результатів та оцінити прискорення;\n"
        "5. Зробити висновок щодо відповідності практичних результатів теоретичній обчислювальній складності O(N²) та O(N log N)."
    )
    r_t.font.name = "Times New Roman"
    r_t.font.size = Pt(14)

    # Код програми
    p_code_title = doc.add_paragraph()
    p_code_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_code_title.paragraph_format.space_before = Pt(10)
    p_code_title.paragraph_format.space_after = Pt(6)
    r_ct = p_code_title.add_run("КОД ПРОГРАМИ")
    r_ct.font.name = "Times New Roman"
    r_ct.font.size = Pt(16)
    r_ct.bold = True

    for line in code_lines:
        p_l = doc.add_paragraph()
        add_code_line(p_l, line)

    # Результати тестування
    doc.add_page_break()
    p_res_title = doc.add_paragraph()
    p_res_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_res_title.paragraph_format.space_before = Pt(10)
    p_res_title.paragraph_format.space_after = Pt(6)
    r_rt = p_res_title.add_run("РЕЗУЛЬТАТИ ТЕСТУВАННЯ ТА РОБОТИ ПРОГРАМИ")
    r_rt.font.name = "Times New Roman"
    r_rt.font.size = Pt(16)
    r_rt.bold = True

    p_test_desc = doc.add_paragraph()
    p_test_desc.paragraph_format.space_before = Pt(2)
    p_test_desc.paragraph_format.space_after = Pt(4)
    r_td = p_test_desc.add_run("1. Результати виконання вбудованих автоматичних тестів (assert) при запуску з прапорцем --test:")
    r_td.font.name = "Times New Roman"
    r_td.font.size = Pt(14)

    test_logs = [
        "> g++ -std=c++20 -O2 main.cpp -o main.exe",
        "> ./main.exe --test",
        "Автотести успішно пройдено."
    ]
    for tl in test_logs:
        p_tl = doc.add_paragraph()
        p_tl.paragraph_format.space_before = Pt(0)
        p_tl.paragraph_format.space_after = Pt(0)
        set_para_shading(p_tl, "1F1F1F")
        r = p_tl.add_run(tl)
        r.font.name = "Consolas"
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(0x4E, 0xC9, 0xB0) if "успішно" in tl else RGBColor(0xDC, 0xDC, 0xDC)

    p_run_desc = doc.add_paragraph()
    p_run_desc.paragraph_format.space_before = Pt(8)
    p_run_desc.paragraph_format.space_after = Pt(4)
    r_rd = p_run_desc.add_run("2. Порівняльна таблиця вимірювань часу виконання алгоритмів (в мілісекундах):")
    r_rd.font.name = "Times New Roman"
    r_rd.font.size = Pt(14)

    bench_logs = [
        "=========================================================================",
        " Порівняння часу виконання: Сортування вставкою vs Двійкове дерево",
        "=========================================================================",
        " |   Розмір |    Тип даних |   Вставкою (мс) |   Деревом (мс) |  Прискорення |",
        " |----------|--------------|-----------------|----------------|--------------|",
        " |     1000 |    Випадкові |           0.074 |          0.120 |        0.620 |",
        " |     5000 |    Випадкові |           1.711 |          0.633 |        2.702 |",
        " |    10000 |    Випадкові |           8.428 |          1.850 |        4.555 |",
        " |    25000 |    Випадкові |          50.158 |          3.777 |       13.280 |",
        " |    50000 |    Випадкові |         182.669 |          8.769 |       20.831 |",
        " |   100000 |    Випадкові |         661.672 |         20.577 |       32.156 |",
        " |----------|--------------|-----------------|----------------|--------------|",
        " |     1000 | Відсортовані |           0.001 |          1.497 |        0.000 |",
        " |     5000 | Відсортовані |           0.003 |         50.104 |        0.000 |",
        " |    10000 | Відсортовані |           0.006 |        202.794 |        0.000 |",
        " |----------|--------------|-----------------|----------------|--------------|",
        " |     1000 |     Зворотні |           0.128 |          1.536 |        0.083 |",
        " |     5000 |     Зворотні |           3.042 |         51.143 |        0.059 |",
        " |    10000 |     Зворотні |          12.676 |        213.268 |        0.059 |",
        "========================================================================="
    ]
    for bl in bench_logs:
        p_bl = doc.add_paragraph()
        p_bl.paragraph_format.space_before = Pt(0)
        p_bl.paragraph_format.space_after = Pt(0)
        set_para_shading(p_bl, "1F1F1F")
        r = p_bl.add_run(bl)
        r.font.name = "Consolas"
        r.font.size = Pt(9.5)
        if "Випадкові" in bl:
            r.font.color.rgb = RGBColor(0x56, 0x9C, 0xD6)
        elif "Відсортовані" in bl:
            r.font.color.rgb = RGBColor(0x4E, 0xC9, 0xB0)
        elif "Зворотні" in bl:
            r.font.color.rgb = RGBColor(0xCE, 0x91, 0x78)
        else:
            r.font.color.rgb = RGBColor(0xDC, 0xDC, 0xDC)

    # Блок-схема (окрема сторінка)
    doc.add_page_break()
    p_flow_title = doc.add_paragraph()
    p_flow_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_flow_title.paragraph_format.space_before = Pt(4)
    p_flow_title.paragraph_format.space_after = Pt(6)
    r_ft = p_flow_title.add_run("БЛОК-СХЕМА АЛГОРИТМУ (UML)")
    r_ft.font.name = "Times New Roman"
    r_ft.font.size = Pt(16)
    r_ft.bold = True

    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(2)
    p_img.paragraph_format.space_after = Pt(2)
    p_img.add_run("[SVG_FLOWCHART_PLACEHOLDER]")

    # Висновок (окрема сторінка)
    doc.add_page_break()
    p_conc = doc.add_paragraph()
    p_conc.paragraph_format.space_before = Pt(12)
    p_conc.paragraph_format.space_after = Pt(6)
    p_conc.paragraph_format.line_spacing = 1.15
    r_c_b = p_conc.add_run("Висновок: ")
    r_c_b.font.name = "Times New Roman"
    r_c_b.font.size = Pt(16)
    r_c_b.bold = True

    r_c_text = p_conc.add_run(
        "У результаті виконання лабораторної роботи було програмно реалізовано та експериментально досліджено "
        "два алгоритми сортування згідно з варіантом 12: сортування вставкою (Insertion sort) та сортування за допомогою "
        "двійкового дерева пошуку (Tree sort).\n\n"
        "На основі аналізу отриманих практичних результатів встановлено:\n"
        "1. На випадкових послідовностях даних сортування двійковим деревом демонструє перевагу над сортуванням вставкою, "
        "яка стрімко зростає зі збільшенням розмірності масиву: для N = 1 000 вставка швидша за рахунок відсутності накладних витрат "
        "на виділення динамічної пам'яті (0.074 мс проти 0.120 мс), проте вже для N = 10 000 сортування деревом у 4.5 рази швидше "
        "(1.850 мс проти 8.428 мс), а для N = 100 000 перевага дерева сягає понад 32 разів (20.577 мс проти 661.672 мс). "
        "Це повністю узгоджується з теоретичними оцінками середньої обчислювальної складності: O(N log N) для Tree sort проти O(N²) для Insertion sort.\n"
        "2. На вже відсортованих послідовностях сортування вставкою досягає свого найкращого випадку O(N) — виконується лише одне "
        "порівняння на кожен елемент і жодного зсуву (0.006 мс для N = 10 000). Натомість просте двійкове дерево пошуку "
        "вироджується в лінійний список довжиною N, що погіршує його складність до найгіршого випадку O(N²) (202.794 мс для N = 10 000).\n"
        "3. На зворотно відсортованих послідовностях обидва алгоритми демонструють квадратичну складність O(N²), проте сортування вставкою "
        "працює на місці (in-place, O(1) додаткової пам'яті), тоді як Tree sort потребує O(N) додаткової пам'яті для вузлів дерева.\n"
        "Таким чином, експериментальні вимірювання повністю підтвердили теоретичні моделі складності обох алгоритмів."
    )
    r_c_text.font.name = "Times New Roman"
    r_c_text.font.size = Pt(14)

    doc.save(docx_path)
    print(f"Попередній DOCX збережено: {docx_path}")

    # Вбудовування векторного SVG через Word COM
    try:
        word = win32com.client.DispatchEx("Word.Application")
        word.Visible = False
        doc_word = word.Documents.Open(docx_path)

        find_range = doc_word.Content
        if find_range.Find.Execute(FindText="[SVG_FLOWCHART_PLACEHOLDER]"):
            find_range.Text = ""
            shape = find_range.InlineShapes.AddPicture(FileName=svg_file)
            shape.LockAspectRatio = -1
            shape.Width = 16.0 * 28.3465
            shape.Height = 23.0 * 28.3465
            print(f"SVG успішно вбудовано у DOCX (ширина: {shape.Width} pt, висота: {shape.Height} pt)")

        doc_word.Save()
        print(f"DOCX оновлено: {docx_path}")

        doc_word.SaveAs(pdf_path, FileFormat=17)
        doc_word.Close()
        word.Quit()
        print(f"PDF успішно створено: {pdf_path}")
    except Exception as e:
        print(f"Помилка Word COM: {e}")

if __name__ == "__main__":
    build_report()
